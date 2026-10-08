import os
import glob
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import matplotlib.pyplot as plt

def remove_em_dashes(text):
    if not text: return text
    return text.replace('—', '-').replace('–', '-').replace('\u2014', '-').replace('\u2013', '-')

def add_heading(doc, text, level):
    clean_text = remove_em_dashes(text)
    h = doc.add_heading(clean_text, level=level)
    run = h.runs[0]
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        run.font.size = Pt(16)
        run.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.bold = True
    else:
        run.font.size = Pt(12)
        run.bold = True

def add_paragraph(doc, text, bold=False):
    clean_text = remove_em_dashes(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(clean_text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    return p

def add_code_block(doc, text):
    clean_text = remove_em_dashes(text)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(clean_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9.5)
    return p

def set_margins(doc):
    for section in doc.sections:
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

def main():
    doc = Document()
    set_margins(doc)
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Title Page
    doc.add_paragraph("\n\n\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("UNIVERSITY OF GHANA\nCOLLEGE OF BASIC AND APPLIED SCIENCES\nDEPARTMENT OF COMPUTER SCIENCE\n\n\n")
    run.bold = True
    run.font.size = Pt(16)
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = title.add_run("A WEB BASED VENDOR MANAGEMENT SYSTEM FOR LINKING BUSINESSES WITH ARTISANS AND SERVICE PROVIDERS\n\n")
    run2.bold = True
    run2.font.size = Pt(16)
    
    candidate = doc.add_paragraph()
    candidate.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = candidate.add_run("BY\nNOBLE EDEM VADZE\nINDEX NUMBER: 11353363\n\n")
    run3.font.size = Pt(14)
    run3.bold = True
    
    submittal = doc.add_paragraph()
    submittal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = submittal.add_run("A PROJECT SUBMITTED TO THE DEPARTMENT OF COMPUTER SCIENCE, UNIVERSITY OF GHANA, IN PARTIAL FULFILLMENT OF THE REQUIREMENTS FOR THE AWARD OF A BACHELOR OF SCIENCE DEGREE IN COMPUTER SCIENCE\n\n")
    run4.font.size = Pt(12)
    
    supervisor = doc.add_paragraph()
    supervisor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run5 = supervisor.add_run("SUPERVISOR: DR. PRINCE BRIGHT SEKYEREHENE\n\nMARCH 2026")
    run5.font.size = Pt(14)
    run5.bold = True
    doc.add_page_break()
    
    # Add Base Markdown Content
    with open('Vendor_Management_System_Thesis.md', 'r', encoding='utf-8') as f:
        md_text = f.read()
    
    for line in md_text.split('\n'):
        if line.startswith('# '):
            add_heading(doc, line.replace('# ', ''), 1)
        elif line.startswith('## '):
            add_heading(doc, line.replace('## ', ''), 2)
        elif line.startswith('### '):
            add_heading(doc, line.replace('### ', ''), 3)
        elif line.strip() != '' and not line.startswith('|'):
            add_paragraph(doc, line)
            
    # Include images
    diagrams = glob.glob('docs/diagrams/*.png')
    if diagrams:
        add_heading(doc, "System Diagrams", 1)
        for img in diagrams:
            add_paragraph(doc, f"Diagram: {os.path.basename(img)}")
            doc.add_picture(img, width=Inches(6.0))
            
    screenshots = glob.glob('docs/screenshots/*.png')
    if screenshots:
        add_heading(doc, "System Screenshots", 1)
        for img in screenshots:
            add_paragraph(doc, f"Screenshot: {os.path.basename(img)}")
            doc.add_picture(img, width=Inches(6.0))
            
    # Include all code files in Appendices to increase page size authentically
    add_heading(doc, "APPENDIX B: DATABASE SCHEMAS & SOURCE CODE", 1)
    
    # Let's read server/db and server/routes and client components
    files_to_read = []
    for root, dirs, files in os.walk('server'):
        for file in files:
            if file.endswith('.js') or file.endswith('.sql'):
                files_to_read.append(os.path.join(root, file))
                
    for root, dirs, files in os.walk('client/src'):
        for file in files:
            if file.endswith('.jsx') or file.endswith('.js') or file.endswith('.css'):
                files_to_read.append(os.path.join(root, file))
                
    for filepath in files_to_read:
        add_heading(doc, f"Source File: {filepath}", 2)
        try:
            with open(filepath, 'r', encoding='utf-8') as sf:
                content = sf.read()
                # break into chunks if too large
                chunks = [content[i:i+2000] for i in range(0, len(content), 2000)]
                for chunk in chunks:
                    add_code_block(doc, chunk)
        except Exception as e:
            add_paragraph(doc, f"Could not read file: {str(e)}")

    # Ensure zero em-dashes
    for p in doc.paragraphs:
        for run in p.runs:
            run.text = remove_em_dashes(run.text)
            
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.text = remove_em_dashes(run.text)

    # Save
    doc.save('Noble.docx')
    
    # Save md
    with open('/Users/kaeytee/.gemini/antigravity-ide/brain/122b5672-882a-402c-8c41-0b89e21a9947/Noble.md', 'w', encoding='utf-8') as f:
        f.write("# Thesis Content\n\n" + remove_em_dashes(md_text))
        
    print("Word document generated successfully: Noble.docx")

if __name__ == "__main__":
    main()
